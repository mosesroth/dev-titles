"""Reference-based audit + review packaging (promoted out of the /tmp scratch).

One home for: a segmentation-independent word-level diff between our output and a
reference subtitle, an LLM classifier of each divergence, and a verdict-labeled /
blind-A/B review sheet. Used to validate transcription (ours vs a human caption)
and translation (ours vs a human translation).

Classifier calibration (operator-marked): `ERROR` (ours wrong) is
reliable and may auto-resolve; `A_OK` (ours right, reference wrong) is OVERCONFIDENT
unless unambiguous, so the prompt keeps a HIGH bar for it and DEFAULTS to `EDGE`;
`A_OK`+`EDGE` both go to a human. The reference is STRONG but NOT gold (a caption may
be auto-generated / spell out identifiers).
"""

from __future__ import annotations

import difflib
import json
import re

# Function words + fillers: a diff that is ONLY these is a style/noise diff, not a
# substantive one (contraction/article/synonym/filler churn between two edited texts).
_FUNCTION = frozenset(
    (
        "a an the this that these those it its there here is are was were be been being am "
        "s re ll m ve d t theres heres thats whats hes shes im youre youll youd weve ive "
        "do does did will would can could shall should may might must have has had we you i "
        "they he she him her them us my your our their of in on at to into onto by with for "
        "from and or but so nor yet as if then than just now some any no not who which what "
        "when where how"
    ).split()
)
_FILLER = frozenset(
    "like essentially basically actually really well right okay ok yeah "
    "kind sort mean know uh um".split()
)


def _key(t: str) -> str:
    return re.sub(r"[^a-z0-9]", "", t.lower())


def _fmt(sec: float) -> str:
    sec = max(0, int(sec))
    return f"{sec // 60:02d}:{sec % 60:02d}"


def _cue_only(span: str) -> bool:
    return all(re.fullmatch(r"[\[\(].*?[\]\)]|♪+", t) for t in span.split())


def _trivial(a: str, b: str) -> bool:
    ig = _FUNCTION | _FILLER
    ca = [k for w in a.split() if (k := _key(w)) and k not in ig]
    cb = [k for w in b.split() if (k := _key(w)) and k not in ig]
    return not ca and not cb


def _tokens(sub):
    """(display, match-key, cue_start) per original word (split on whitespace+hyphens,
    not apostrophes) — display keeps real orthography, key normalizes for matching."""
    out = []
    for c in sub.cues:
        for t in re.split(r"[\s\-]+", " ".join(c.lines)):
            if _key(t):
                out.append((t, _key(t), c.start))
    return out


def _content_keys(keys):
    ig = _FUNCTION | _FILLER
    return [k for k in keys if k and k not in ig]


def _present_elsewhere(span_keys, opposite_keyset, *, frac: float = 0.8) -> bool:
    """True if a MULTI-WORD content span's keys are (mostly) present in the opposite
    full token stream — i.e. the clause exists, just at a different position. difflib
    is order-sensitive, so a clause our CJK re-segmentation RE-TIMED shows up as a bogus
    insert/delete ('(nothing)') even though the content is there. Guards against calling
    that a real omission/addition."""
    uniq = set(_content_keys(span_keys))
    if len(uniq) < 2:  # a single content word can coincide by chance
        return False
    return sum(1 for k in uniq if k in opposite_keyset) / len(uniq) >= frac


def word_diffs(ours_sub, ref_sub, *, ctx: int = 7) -> list[dict]:
    """Substantive, timestamped word-level differences between OURS and a REFERENCE
    subtitle — robust to different segmentation. Skips cue-only, function/filler-only
    (trivial), and RE-TIMING artifacts (a clause present elsewhere in the opposite
    stream — not a real omission/addition). Returns [{ts, ours, ref, ctx}]."""
    A, B = _tokens(ours_sub), _tokens(ref_sub)
    Ad, Ak, At = [x[0] for x in A], [x[1] for x in A], [x[2] for x in A]
    Bd, Bk = [x[0] for x in B], [x[1] for x in B]
    ours_keyset, ref_keyset = set(Ak), set(Bk)
    rows = []
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, Ak, Bk, autojunk=False).get_opcodes():
        if tag == "equal":
            continue
        a = " ".join(Ad[i1:i2]).strip()
        b = " ".join(Bd[j1:j2]).strip()
        if (_cue_only(a) and _cue_only(b)) or _trivial(a, b):
            continue
        # Re-timing filter: a pure ref-only span (ours "(nothing)") whose content is
        # present elsewhere in OURS — or a pure ours-only span present elsewhere in REF —
        # is re-segmentation, not a real omission/addition.
        if tag == "insert" and _present_elsewhere(Bk[j1:j2], ours_keyset):
            continue
        if tag == "delete" and _present_elsewhere(Ak[i1:i2], ref_keyset):
            continue
        ts = At[i1] if i1 < i2 else (At[i1 - 1] if i1 > 0 else 0)
        rows.append(
            {
                "ts": _fmt(ts),
                "ours": a or "(nothing)",
                "ref": b or "(nothing)",
                "ctx": " ".join(Ad[max(0, i1 - ctx) : i1]) + " ⟦?⟧ " + " ".join(Ad[i2 : i2 + ctx]),
            }
        )
    return rows


_TX_PROMPT = (
    "Audit an automatic transcription (OURS) of a developer video against a HUMAN caption "
    "that is a STRONG but NOT perfect reference (it may be auto-generated or spell out code "
    "identifiers). Judge by what was plausibly SAID. DEFAULT to EDGE; only leave EDGE when a "
    "rule clearly decides. For each item:\n"
    "- ERROR: OURS is clearly wrong (misheard / dropped / added / mangled identifier).\n"
    "- A_OK: OURS is UNAMBIGUOUSLY correct and the reference is wrong — a verbatim code "
    "identifier, a DCMP number, a correct name, or clearly-spoken content the caption dropped. "
    "Use SPARINGLY; if unsure, EDGE.\n"
    "- EDGE: ambiguous / equally valid / needs the audio.\n"
)
_TL_PROMPT = (
    "Audit our {lang} subtitle translation (OURS) against a HUMAN translation (a strong "
    "reference, not the only valid rendering). For each item:\n"
    "- ERROR: OURS is defective vs the SOURCE — mistranslation, omission, wrong/inconsistent "
    "term, register (ja not です・ます; zh not 您/formal), or a term split badly.\n"
    "- A_OK: OURS is an equally acceptable rendering (most differences are this — do not flag "
    "mere paraphrase).\n"
    "- EDGE: genuinely unsure / a nuance for a native reviewer.\n"
    "IMPORTANT: when OURS is '(nothing)', READ THE ctx — our re-segmentation re-times/re-orders "
    "cues and localizes terms the reference kept in Latin (Developer Options vs 开发者选项), so a "
    "clause the diff aligned to '(nothing)' is OFTEN present in the OURS context or elsewhere. "
    "Only mark ERROR if the meaning is GENUINELY ABSENT from OURS; if the ctx already conveys it "
    "(re-timed or localized differently), it is A_OK.\n"
)


def classify(
    rows: list[dict], invoke, *, domain: str = "transcription", lang: str = "", batch: int = 18
) -> list[dict]:
    """LLM-classify each diff row (adds 'verdict' in ERROR/A_OK/EDGE + 'why'). Batched
    and guarded so one failed call doesn't lose the run."""
    lead = _TL_PROMPT.format(lang=lang or "target") if domain == "translation" else _TX_PROMPT
    for off in range(0, len(rows), batch):
        chunk = rows[off : off + batch]
        listing = "\n".join(
            f"{k+1}. OURS: {r['ours']!r} | REF: {r['ref']!r} | ctx: …{r['ctx']}…"
            for k, r in enumerate(chunk)
        )
        prompt = (
            lead + "Output ONLY a JSON array: "
            '{"n":N,"verdict":"ERROR|A_OK|EDGE","why":"<=8 words"}.\n\n' + listing
        )
        try:
            raw = invoke(prompt) or ""
            s, e = raw.find("["), raw.rfind("]")
            by = {d["n"]: d for d in json.loads(raw[s : e + 1]) if isinstance(d, dict)}
        except Exception:
            by = {}
        for k, r in enumerate(chunk):
            v = by.get(k + 1) or {}
            r["verdict"] = v.get("verdict", "EDGE")
            r["why"] = v.get("why", "")
    return rows


def write_review_xlsx(rows: list[dict], path, title: str = "") -> dict:
    """Verdict-labeled review sheet (ERROR/EDGE first; A_OK below). 'Your call' column
    for the reviewer. Returns the verdict counts. ERROR auto-resolves; A_OK+EDGE = human."""
    from collections import Counter

    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    order = {"ERROR": 0, "EDGE": 1, "A_OK": 2}
    rows = sorted(rows, key=lambda r: order.get(r.get("verdict", ""), 3))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "audit"
    ws.append(
        [
            f"{title} — OURS vs reference. My verdict is a hint; put the correct reading in "
            "'Your call'. ERROR=ours wrong, A_OK=ours right, EDGE=toss-up. Scrub to Time."
        ]
    )
    ws.append(
        [
            "My verdict",
            "Time",
            "Ours",
            "Reference",
            "Why",
            "Context",
            "Your call (ours / ref / both wrong)",
        ]
    )
    for c in ws[2]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    fills = {"ERROR": "FFC7CE", "EDGE": "FFF2CC", "A_OK": "C6EFCE"}
    for r in rows:
        ws.append(
            [
                r.get("verdict", "?"),
                r.get("ts", ""),
                r["ours"],
                r["ref"],
                r.get("why", ""),
                "…" + r["ctx"] + "…",
                "",
            ]
        )
        ws.cell(ws.max_row, 1).fill = PatternFill(
            "solid", fgColor=fills.get(r.get("verdict", ""), "FFFFFF")
        )
        ws.cell(ws.max_row, 7).fill = PatternFill("solid", fgColor="FFF2CC")
    for col, w in zip("ABCDEFG", (11, 7, 24, 24, 24, 48, 26)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=3):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A3"
    from pathlib import Path

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return dict(Counter(r.get("verdict", "?") for r in rows))
