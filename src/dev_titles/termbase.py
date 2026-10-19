"""Import a term-base export (.xlsx) with en_us/ja_jp/zh_cn columns -> glossary/termbase.yaml.

A term base (TB) is a localization vendor's AUTHORITATIVE multilingual term list (typically
hundreds to thousands of approved concepts). We import EN source + ja_jp + zh_cn renderings (any `forbidden`-flagged variant
becomes an avoid-form) as SOFT glossary terms: injected as the preferred rendering + used for
cross-cue consistency, NOT hard-substituted (thousands of substring substitutions would risk
corrupting unrelated words). `glossary.load` merges termbase.yaml, with the hand-curated
glossary.yaml WINNING on any shared source form.

Regenerate when the TB updates:
    python -m dev_titles.termbase "<termbase.xlsx>" glossary/termbase.yaml
"""

from __future__ import annotations

from pathlib import Path

# field offsets from a language-value column (header block: TID, <lang>, status, forbidden,
# case, exact, note, usage, created_by, created_at, modified_by, modified_at)
_FORBIDDEN = 2


def import_termbase(xlsx_path: str | Path) -> list[dict]:
    """Parse the term-base workbook into a list of glossary-term dicts (deduped by source)."""
    try:
        import openpyxl
    except ImportError as e:  # pragma: no cover
        raise RuntimeError("Reading .xlsx term bases needs: pip install 'dev-titles[xlsx]'") from e

    ws = openpyxl.load_workbook(xlsx_path)["terms"]
    hdr = [str(c.value) if c.value is not None else "" for c in ws[1]]

    def cols(lang: str) -> list[int]:
        return [i for i, h in enumerate(hdr) if h == lang]

    rows = ws.iter_rows(min_row=2, values_only=True)

    en_c, ja_c, zh_c = cols("en_us"), cols("ja_jp"), cols("zh_cn")

    def collect(row, cs) -> tuple[list[str], list[str]]:
        ok: list[str] = []
        bad: list[str] = []
        for c in cs:
            v = row[c]
            if not v or not str(v).strip():
                continue
            forbidden = str(row[c + _FORBIDDEN]).lower() == "true"
            (bad if forbidden else ok).append(str(v).strip())
        return ok, bad

    out: list[dict] = []
    seen: set[str] = set()
    for row in rows:
        en, _ = collect(row, en_c)
        ja, ja_bad = collect(row, ja_c)
        zh, zh_bad = collect(row, zh_c)
        if not en or not (ja or zh):
            continue
        key = en[0].lower()
        if key in seen:
            continue
        seen.add(key)
        t: dict = {"source": en[0]}
        extra = [a for a in en[1:] if a.lower() != key]
        if extra:
            t["aliases"] = extra
        if zh:
            t["zh-Hans"] = zh[0]
        if ja:
            t["ja"] = ja[0]
        avoid = {}
        if zh_bad:
            avoid["zh-Hans"] = zh_bad
        if ja_bad:
            avoid["ja"] = ja_bad
        if avoid:
            t["avoid"] = avoid
        t["enforce"] = "soft"
        out.append(t)
    return out


def write_termbase_yaml(terms: list[dict], path: str | Path) -> None:
    import yaml

    body = yaml.safe_dump(
        {"languages": ["zh-Hans", "ja"], "terms": terms},
        allow_unicode=True,
        sort_keys=False,
        width=1000,
    )
    Path(path).write_text(
        "# GENERATED from the term-base export — DO NOT hand-edit (glossary.yaml wins on\n"
        "# shared source forms). Regenerate: python -m dev_titles.termbase <TB.xlsx> "
        + str(path)
        + "\n"
        + body,
        encoding="utf-8",
    )


if __name__ == "__main__":
    import sys

    src, dst = sys.argv[1], sys.argv[2]
    terms = import_termbase(src)
    write_termbase_yaml(terms, dst)
    print(f"wrote {len(terms)} terms -> {dst}")
